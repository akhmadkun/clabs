import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
DOCS=ROOT/'docs'
TOPO=ROOT/'topology'/'SYS-101_TOPOLOGY.png'
NAVY='102A43'; BLUE='1F6FB2'; LIGHT='EAF2F8'; GOLD='FFF2CC'; GREEN='E2F0D9'; RED='FCE4D6'; GRAY='52606D'; CODE='0B1F33'

def shade(cell,color):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),color)

def set_cell_margins(cell,top=80,start=120,bottom=80,end=120):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def set_table_widths(table,widths):
    table.autofit=False; table.alignment=WD_TABLE_ALIGNMENT.LEFT
    tblPr=table._tbl.tblPr
    tblW=tblPr.find(qn('w:tblW'))
    if tblW is None: tblW=OxmlElement('w:tblW'); tblPr.append(tblW)
    tblW.set(qn('w:w'),str(sum(widths))); tblW.set(qn('w:type'),'dxa')
    tblInd=tblPr.find(qn('w:tblInd'))
    if tblInd is None: tblInd=OxmlElement('w:tblInd'); tblPr.append(tblInd)
    tblInd.set(qn('w:w'),'120'); tblInd.set(qn('w:type'),'dxa')
    grid=table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        gc=OxmlElement('w:gridCol'); gc.set(qn('w:w'),str(w)); grid.append(gc)
    for row in table.rows:
        for cell,w in zip(row.cells,widths):
            tcPr=cell._tc.get_or_add_tcPr(); tcW=tcPr.find(qn('w:tcW'))
            if tcW is None: tcW=OxmlElement('w:tcW'); tcPr.append(tcW)
            tcW.set(qn('w:w'),str(w)); tcW.set(qn('w:type'),'dxa'); set_cell_margins(cell)

def set_font(run,name='Aptos',size=10.5,color=None,bold=None,italic=None):
    run.font.name=name; run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),name); run._element.rPr.rFonts.set(qn('w:hAnsi'),name)
    run.font.size=Pt(size)
    if color: run.font.color.rgb=RGBColor.from_string(color)
    if bold is not None: run.bold=bold
    if italic is not None: run.italic=italic

def add_inline(p,text,size=10.5,color=None,bold=False,italic=False,font='Aptos'):
    parts=re.split(r'(\*\*[^*]+\*\*|`[^`]+`)',text)
    for part in parts:
        if not part: continue
        if part.startswith('**'):
            r=p.add_run(part[2:-2]); set_font(r,font,size,color,True,italic)
        elif part.startswith('`'):
            r=p.add_run(part[1:-1]); set_font(r,'Consolas',size-0.5,'9C0006',False,italic)
        else:
            r=p.add_run(part); set_font(r,font,size,color,bold,italic)

def configure_doc(doc,title,kind):
    sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=sec.footer_distance=Inches(.492)
    styles=doc.styles
    normal=styles['Normal']; normal.font.name='Aptos'; normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string('1F2937')
    normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.15
    for name,size,color,before,after in [('Heading 1',16,BLUE,18,10),('Heading 2',13,BLUE,14,7),('Heading 3',11.5,NAVY,10,5)]:
        st=styles[name]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
        st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
    header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run(f'SYS-101 | {kind}'),'Aptos',8,GRAY,True)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run('JNCIE-SP Modular Lab Series | SYS-101'),'Aptos',8,GRAY)
    # cover
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(72); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('JNCIE-SP MODULAR LAB'); set_font(r,'Aptos',12,BLUE,True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(8)
    r=p.add_run('SYS-101'); set_font(r,'Aptos Display',30,NAVY,True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('Initial System Settings'); set_font(r,'Aptos Display',22,BLUE,True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(18)
    r=p.add_run(kind); set_font(r,'Aptos',14,GRAY,False,True)
    if TOPO.exists(): doc.add_picture(str(TOPO),width=Inches(6.25)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    t=doc.add_table(rows=1,cols=3); set_table_widths(t,[3120,3120,3120]); shade(t.cell(0,0),LIGHT); shade(t.cell(0,1),LIGHT); shade(t.cell(0,2),LIGHT)
    vals=[('PLATFORM','vJunos-router 25.4'),('DURATION','3 hours'),('TARGET','2 repetitions')]
    for cell,(lab,val) in zip(t.rows[0].cells,vals):
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(lab+'\n'),'Aptos',8,GRAY,True); set_font(p.add_run(val),'Aptos',10,NAVY,True)
    doc.add_page_break()

def add_callout(doc,label):
    colors={'Objective':LIGHT,'Important':GOLD,'Explanation':LIGHT,'Expected Result':GREEN,'Required evidence':RED,'Verification':LIGHT}
    t=doc.add_table(rows=1,cols=1); set_table_widths(t,[9360]); shade(t.cell(0,0),colors.get(label,LIGHT))
    p=t.cell(0,0).paragraphs[0]; set_font(p.add_run(label),'Aptos',10,NAVY,True)

def build(md_path,out_path,kind):
    lines=md_path.read_text().splitlines(); doc=Document(); configure_doc(doc,'SYS-101',kind)
    in_code=False; code=[]; i=0; first_h1=True
    while i<len(lines):
        line=lines[i]
        if line.startswith('```'):
            if not in_code: in_code=True; code=[]
            else:
                t=doc.add_table(rows=1,cols=1); set_table_widths(t,[9360]); shade(t.cell(0,0),CODE)
                p=t.cell(0,0).paragraphs[0]; p.paragraph_format.space_after=Pt(0)
                for j,c in enumerate(code):
                    r=p.add_run(c+('\n' if j<len(code)-1 else '')); set_font(r,'Consolas',8.5,'FFFFFF')
                in_code=False
            i+=1; continue
        if in_code: code.append(line); i+=1; continue
        if not line.strip(): i+=1; continue
        if line.startswith('# '): first_h1=False; i+=1; continue
        if line.startswith('## '): doc.add_paragraph(line[3:],style='Heading 1'); i+=1; continue
        if line.startswith('### '):
            lab=line[4:]
            if lab in {'Objective','Important','Explanation','Expected Result','Required evidence','Verification'}: add_callout(doc,lab)
            else: doc.add_paragraph(lab,style='Heading 2')
            i+=1; continue
        if line.startswith('|'):
            tbl=[]
            while i<len(lines) and lines[i].startswith('|'):
                vals=[x.strip() for x in lines[i].strip('|').split('|')]
                if not all(set(x)<=set('-: ') for x in vals): tbl.append(vals)
                i+=1
            if tbl:
                cols=len(tbl[0]); t=doc.add_table(rows=len(tbl),cols=cols); widths=[int(9360/cols)]*cols; widths[-1]+=9360-sum(widths); set_table_widths(t,widths)
                for rr,row in enumerate(tbl):
                    for cc,val in enumerate(row):
                        cell=t.cell(rr,cc); cell.text=''; p=cell.paragraphs[0]; add_inline(p,val,9.2,bold=(rr==0))
                        if rr==0: shade(cell,BLUE); [setattr(run.font.color,'rgb',RGBColor(255,255,255)) for run in p.runs]
            continue
        m=re.match(r'^(\d+)\.\s+(.*)',line)
        if m:
            p=doc.add_paragraph(style='List Number'); add_inline(p,m.group(2)); i+=1; continue
        if line.startswith('- '):
            p=doc.add_paragraph(style='List Bullet'); add_inline(p,line[2:]); i+=1; continue
        p=doc.add_paragraph(); add_inline(p,line)
        i+=1
    doc.save(out_path)

build(DOCS/'SYS-101_TASK-ONLY.md',DOCS/'SYS-101_TASK-ONLY.docx','Task-Only Assessment')
build(DOCS/'SYS-101_SOLUTION-GUIDE.md',DOCS/'SYS-101_SOLUTION-GUIDE.docx','Task and Solution Guide')
print('documents built')

